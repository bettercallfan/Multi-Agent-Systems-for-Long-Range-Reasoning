import json
from pathlib import Path


def read_docx(path: str):
    result = {
        "path": path,
        "type": "docx",
        "status": "unknown",
        "text_preview": "",
        "paragraph_count": 0,
        "tables_count": 0,
        "error": "",
    }

    try:
        from docx import Document

        doc = Document(path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        result["status"] = "success"
        result["paragraph_count"] = len(paragraphs)
        result["tables_count"] = len(doc.tables)
        result["text_preview"] = "\n".join(paragraphs[:80])[:6000]

    except Exception as e:
        result["status"] = "error"
        result["error"] = str(e)

    return result


def read_xlsx(path: str):
    result = {
        "path": path,
        "type": "xlsx",
        "status": "unknown",
        "sheets": [],
        "sheet_summaries": {},
        "error": "",
    }

    try:
        import pandas as pd

        with pd.ExcelFile(path, engine="openpyxl") as excel:
            result["sheets"] = excel.sheet_names

            for sheet in excel.sheet_names:
                df = pd.read_excel(excel, sheet_name=sheet)

                sample = df.head(5).fillna("")
                # Pandas 3.x 兼容：优先 map，fallback applymap
                def _safe_convert(x):
                    if hasattr(x, "isoformat"):
                        return x.isoformat()
                    return x

                try:
                    sample = sample.map(_safe_convert)
                except AttributeError:
                    sample = sample.applymap(_safe_convert)

                result["sheet_summaries"][sheet] = {
                    "rows": int(len(df)),
                    "columns": [str(c) for c in df.columns],
                    "dtypes": {str(k): str(v) for k, v in df.dtypes.items()},
                    "sample_rows": sample.to_dict(orient="records"),
                }

        result["status"] = "success"

    except Exception as e:
        result["status"] = "error"
        result["error"] = str(e)

    return result


def read_pdf(path: str):
    result = {
        "path": path,
        "type": "pdf",
        "status": "unknown",
        "page_count": 0,
        "text_preview": "",
        "error": "",
    }

    try:
        text_parts = []

        try:
            import pdfplumber

            with pdfplumber.open(path) as pdf:
                result["page_count"] = len(pdf.pages)

                for page in pdf.pages[:10]:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)

        except Exception:
            import PyPDF2

            with open(path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                result["page_count"] = len(reader.pages)

                for page in reader.pages[:10]:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)

        result["status"] = "success"
        result["text_preview"] = "\n".join(text_parts)[:10000]

    except Exception as e:
        result["status"] = "error"
        result["error"] = str(e)

    return result


def read_file_preview(path: str):
    suffix = Path(path).suffix.lower()

    if suffix == ".docx":
        return read_docx(path)

    if suffix in [".xlsx", ".xls"]:
        return read_xlsx(path)

    if suffix == ".pdf":
        return read_pdf(path)

    result = {
        "path": path,
        "type": suffix.replace(".", "") or "unknown",
        "status": "unknown",
        "text_preview": "",
        "error": "",
    }

    try:
        # Preview only a bounded prefix. Large JSONL/CSV files and binary
        # captures must never be loaded wholesale merely for classification.
        with Path(path).open("rb") as stream:
            prefix = stream.read(20_000)
        text = prefix.decode("utf-8-sig", errors="ignore")
        result["status"] = "success"
        result["text_preview"] = text[:5000]
    except Exception as e:
        result["status"] = "error"
        result["error"] = str(e)

    return result


def build_file_previews(files):
    previews = []
    paths = [Path(file_path) for file_path in files]
    json_paths = [path for path in paths if path.suffix.lower() == ".json"]
    aggregated_json: set[Path] = set()
    if len(json_paths) > 100:
        sample_paths = sorted(json_paths)[:3]
        sample_previews = [read_file_preview(str(path)) for path in sample_paths]
        previews.append({
            "path": str(json_paths[0].parent / "*.json"),
            "type": "json_collection",
            "status": "success",
            "file_count": len(json_paths),
            "sample_files": [path.name for path in sample_paths],
            "text_preview": "\n".join(
                item.get("text_preview", "")[:1500] for item in sample_previews
            )[:5000],
            "error": "",
        })
        aggregated_json = set(json_paths)

    for path in paths:
        if path in aggregated_json:
            continue
        previews.append(read_file_preview(str(path)))

    return previews


def save_file_previews(previews, output_path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    output_path.write_text(
        json.dumps(previews, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
