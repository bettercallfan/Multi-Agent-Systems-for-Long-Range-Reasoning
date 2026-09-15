import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, patch
from types import SimpleNamespace

from autogen_ext.models.replay import ReplayChatCompletionClient

from main import apply_runtime_policy_overrides
from agents.code_agents import execute_code_file, save_generated_code, snapshot_run_files
from orchestration.task.artifact_validator import validate_intermediate_artifacts
from orchestration.task.input_preparation import prepare_normalized_input
from orchestration.core.run_state import RunState
from orchestration.core.schemas import CodeGenerationResult, ExecutionResult
from orchestration.task.task_normalizer import finalize_task_spec, normalize_to_task_spec
from orchestration.core.workflow import (
    _build_verified_report_fallback,
    _finish_planning_blocked,
    run_explicit_workflow,
)
from orchestration.core.prompt_builder import build_code_prompt
from orchestration.execution.default_executors import (
    CodePipelineNodeExecutor,
    _failed_artifact_diagnostics,
)
from orchestration.execution.node_executor import NodeExecutionContext
from orchestration.graph.task_graph import TaskNode
from utils.output import save_agent_trace
from utils.review_artifacts import final_validation, pre_report_review


class TaskSpecTests(unittest.TestCase):
    def test_verified_report_fallback_contains_required_sections_and_evidence(self):
        sections = [
            "任务理解", "任务拆解", "方法或模型设计",
            "代码实现或执行过程", "结果分析", "风险与改进建议",
        ]
        report = _build_verified_report_fallback(
            {
                "task_name": "装箱优化",
                "task_type": "math_modeling",
                "required_report_sections": sections,
            },
            {
                "task_graph": {"nodes": [{
                    "node_id": "solve", "description": "求解装箱方案",
                    "capability": "code", "status": "completed",
                }]},
                "execution": {"exit_code": 0, "phase": "execution"},
                "business_validation": {"status": "passed"},
                "artifacts": {"artifacts/result.json": {
                    "vehicle_count": 6, "total_cost": 4200,
                }},
            },
        )
        self.assertGreater(len(report), 200)
        for section in sections:
            self.assertIn(f"## {section}", report)
        self.assertIn("artifacts/result.json", report)
        self.assertNotIn("ReportAgent", report)

    def test_complex_code_policy_allows_real_multistage_implementation(self):
        spec = normalize_to_task_spec({
            "files": [],
            "text": "建立优化模型并编写求解算法代码，执行仿真计算",
        })
        self.assertEqual(spec["code_policy"]["mode"], "complex")
        self.assertEqual(spec["code_policy"]["max_lines"], 320)

    def test_input_preparation_is_task_independent(self):
        prepared = prepare_normalized_input({
            "input": {
                "type": "file",
                "files": ["/temporary/location/source.txt"],
                "text": "分析输入",
            },
            "file_previews": [{
                "path": "source.txt", "type": "txt",
                "status": "success", "text_preview": "evidence",
            }],
        })

        self.assertEqual(prepared["input"]["files"], ["source.txt"])
        self.assertEqual(prepared["input"]["text"], "分析输入")
        self.assertEqual(
            prepared["file_previews"][0]["text_preview"], "evidence",
        )

    def test_docx_tables_are_preserved_by_generic_input_boundary(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.docx"
            document = Document()
            document.add_paragraph("通用说明")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "字段A"
            table.cell(0, 1).text = "字段B"
            table.cell(1, 0).text = "值1"
            table.cell(1, 1).text = "值2"
            document.save(path)

            prepared = prepare_normalized_input({
                "input": {
                    "type": "file",
                    "files": [str(path)],
                    "text": "",
                },
                "file_previews": [],
            })

        source = prepared["structured_sources"][0]
        self.assertEqual(source["status"], "loaded")
        self.assertEqual(source["kind"], "document")
        self.assertEqual(source["content"], "通用说明")
        self.assertEqual(
            source["tables"][0]["rows"],
            [["字段A", "字段B"], ["值1", "值2"]],
        )
        self.assertIn(
            "two-dimensional",
            prepared["format_contract"]["table_rows"],
        )

    def test_runtime_llmlingua_options_are_frozen_into_task_spec(self):
        spec = normalize_to_task_spec({"files": [], "text": "分析任务"})
        args = SimpleNamespace(
            llmlingua_device_map="cuda",
            llmlingua_allow_download=True,
        )

        apply_runtime_policy_overrides(spec, args)
        spec = finalize_task_spec(spec, [])

        self.assertEqual(spec["communication_policy"]["llmlingua_device_map"], "cuda")
        self.assertTrue(spec["communication_policy"]["llmlingua_allow_download"])

    def test_generic_task_has_single_authoritative_contract(self):
        raw = {
            "input_type": "text",
            "files": [],
            "text": (
                "请完成文档分析并生成 artifacts/result.json。\n"
                "最终报告至少包含：\n"
                "1. 任务理解\n2. 结果分析\n"
                "不要给出最终决定。"
            ),
        }
        spec = normalize_to_task_spec(raw)
        spec = finalize_task_spec(spec, [])
        self.assertEqual(spec["task_type"], "document_analysis")
        self.assertEqual(spec["code_policy"]["mode"], "lightweight")
        self.assertIn("artifacts/result.json", spec["required_artifacts"])
        self.assertIn("artifacts/code_pipeline.py", spec["required_artifacts"])
        self.assertNotIn("plugin_id", spec)
        self.assertEqual(spec["artifact_contract"]["final_artifacts"], ["final_report.md"])
        self.assertIn("artifacts/result.json", spec["artifact_contract"]["intermediate_artifacts"])
        self.assertIn("run_state.json", spec["artifact_contract"]["framework_artifacts"])
        self.assertEqual(spec["required_report_sections"], ["任务理解", "结果分析"])
        self.assertEqual(spec["report_policy"]["decision_scope"], "descriptive_only")
        self.assertTrue(spec["report_policy"]["forbid_unsupported_decisions"])
        self.assertEqual(spec["routing_policy"]["mode"], "agent_prune_lite")
        self.assertTrue(spec["routing_policy"]["direct_dependencies_only"])
        self.assertEqual(spec["communication_policy"]["compressor"], "auto")
        self.assertEqual(spec["communication_policy"]["max_message_tokens"], 1500)
        self.assertEqual(spec["communication_policy"]["max_node_context_bytes"], 48000)
        self.assertEqual(spec["communication_policy"]["max_control_prompt_tokens"], 8000)
        self.assertEqual(spec["communication_policy"]["max_inline_structured_bytes"], 8000)
        self.assertTrue(spec["communication_policy"]["persist_message_payloads"])
        self.assertEqual(spec["team_policy"]["mode"], "dylan_lite")
        self.assertNotIn("max_executor_switches", spec["team_policy"])
        self.assertEqual(spec["recovery_policy"]["max_replans"], 1)


class RunStateTests(unittest.TestCase):
    def test_mutable_framework_files_do_not_publish_stale_self_hashes(self):
        with tempfile.TemporaryDirectory() as directory:
            state = RunState(directory)
            (Path(directory) / "agent_trace.md").write_text("first", encoding="utf-8")
            (Path(directory) / "artifact.json").write_text("{}", encoding="utf-8")

            state.refresh_artifacts()
            metadata = state.to_dict()["artifacts"]["metadata"]

            self.assertNotIn("run_state.json", metadata)
            self.assertNotIn("agent_trace.md", metadata)
            self.assertIn("artifact.json", metadata)

    def test_artifact_missing_does_not_overwrite_upstream_execution_failure(self):
        with tempfile.TemporaryDirectory() as tmp:
            state = RunState(tmp)
            state.start_stage("execute")
            state.set_failure("execution_failure", "code", "execute", ["连接失败"])
            state.record_artifact_quality({
                "status": "failed",
                "checked_artifacts": ["artifacts/result.json"],
                "issues": ["缺少中间产物"],
                "failure_type": "artifact_missing",
                "repair_target": "code",
                "resume_stage": "execute",
            })
            self.assertEqual(state.get("failure")["failure_type"], "execution_failure")
            self.assertTrue(any(
                event["type"] == "derived_failure_recorded" for event in state.get("events")
            ))

    def test_real_execution_result_controls_state(self):
        with tempfile.TemporaryDirectory() as directory:
            state = RunState(directory, code_policy={"mode": "lightweight", "max_retries": 1})
            state.start_stage("classify")
            state.start_stage("plan")
            state.start_stage("execute")
            state.record_execution(ExecutionResult(
                attempt=1, command=["python", "x.py"], exit_code=1, stderr="boom"
            ))
            self.assertTrue(state.to_dict()["execution"]["failed"])
            self.assertEqual(state.to_dict()["execution"]["exit_code"], 1)


class ExecutorTests(unittest.IsolatedAsyncioTestCase):
    async def test_failed_json_artifact_is_available_as_bounded_repair_evidence(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            (run_dir / "artifacts").mkdir()
            (run_dir / "artifacts" / "result.json").write_text(
                json.dumps({
                    "status": "failed",
                    "error": "float() argument must not be None",
                }),
                encoding="utf-8",
            )
            diagnostics = _failed_artifact_diagnostics(
                run_dir,
                ["artifacts/code_pipeline.py", "artifacts/result.json"],
            )
            self.assertEqual(
                diagnostics["artifacts/result.json"]["status"], "failed",
            )
            self.assertIn(
                "must not be None",
                diagnostics["artifacts/result.json"]["error"],
            )

    async def test_executor_uses_run_dir_without_cd(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            (run_dir / "artifacts").mkdir()
            (run_dir / "artifacts" / "code_pipeline.py").write_text(
                "from pathlib import Path\n"
                "Path('artifacts/result.json').write_text('{\"ok\": true}', encoding='utf-8')\n",
                encoding="utf-8",
            )
            result = await execute_code_file(run_dir, "artifacts/code_pipeline.py", attempt=1)
            self.assertEqual(result.exit_code, 0)
            self.assertIn("artifacts/result.json", result.produced_files)
            self.assertNotIn("artifacts/code_pipeline.py", result.produced_files)

    async def test_execution_reports_only_new_or_modified_files(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            (run_dir / "artifacts").mkdir()
            (run_dir / "run_state.json").write_text('{"stage":"execute"}', encoding="utf-8")
            (run_dir / "artifacts" / "result.json").write_text('{"value": 1}', encoding="utf-8")
            script = run_dir / "artifacts" / "code_pipeline.py"
            before_generation = snapshot_run_files(run_dir)
            script.write_text(
                "from pathlib import Path\n"
                "Path('artifacts/result.json').write_text('{\"value\": 2}', encoding='utf-8')\n",
                encoding="utf-8",
            )

            result = await execute_code_file(
                run_dir,
                "artifacts/code_pipeline.py",
                attempt=1,
                before_snapshot=before_generation,
            )

            self.assertEqual(result.exit_code, 0)
            self.assertEqual(result.produced_files, [
                "artifacts/code_pipeline.py", "artifacts/result.json",
            ])
            self.assertNotIn("run_state.json", result.produced_files)

    async def test_execution_rejects_script_that_rewrites_its_own_entrypoint(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            script = run_dir / "artifacts" / "code_pipeline.py"
            script.parent.mkdir(parents=True)
            script.write_text(
                "from pathlib import Path\n"
                "target = Path('artifacts') / ('code_' + 'pipeline.py')\n"
                "target.write_text('.version = \\\"1.0\\\"\\n', encoding='utf-8')\n"
                "Path('artifacts/result.json').write_text('{}', encoding='utf-8')\n",
                encoding="utf-8",
            )

            result = await execute_code_file(
                run_dir, "artifacts/code_pipeline.py", attempt=1,
            )

            self.assertEqual(result.exit_code, 66)
            self.assertIn("修改了自身入口文件", result.stderr)


class PromptContractTests(unittest.TestCase):
    def test_code_prompt_contains_machine_acceptance_obligations(self):
        task_spec = {
            "task_id": "contract-run",
            "task_name": "generic computation",
            "task_type": "math_modeling",
            "input": {"type": "directory", "files": [], "text": ""},
            "code_policy": {"max_lines": 200},
            "artifact_contract": {
                "intermediate_artifacts": [
                    "artifacts/result.json", "artifacts/code_pipeline.py",
                ],
            },
            "required_artifacts": [
                "artifacts/result.json", "artifacts/code_pipeline.py",
            ],
            "artifacts": {"code": "artifacts/code_pipeline.py"},
            "requirement_contract": {
                "requirements": [{
                    "requirement_id": "REQ-RESULT",
                    "statement": "输出全部记录",
                    "mandatory": True,
                    "acceptance_criteria": [{
                        "criterion_id": "AC-RESULT",
                        "method": "required_fields",
                        "target": "artifacts/result.json",
                        "params": {"fields": ["/records"]},
                        "condition": "包含记录",
                        "severity": "blocking",
                    }],
                }],
            },
        }

        prompt = build_code_prompt(task_spec, {"current_node": "code"})

        self.assertIn("AC-RESULT", prompt)
        self.assertIn('"/records"', prompt)
        self.assertIn("不能弱化", prompt)
        self.assertIn("不得依赖运行时 pip install", prompt)
        self.assertIn("__previous_code_candidate__", prompt)


class CodePersistenceTests(unittest.TestCase):
    def test_save_generated_code_accepts_relative_run_dir(self):
        with tempfile.TemporaryDirectory(dir=".") as directory:
            run_dir = Path(directory).resolve()
            result = CodeGenerationResult(
                code="from pathlib import Path\nPath('artifacts/result.json').write_text('{}')",
                entrypoint="artifacts/code_pipeline.py",
                expected_outputs=["artifacts/result.json"],
            )
            relative_run_dir = run_dir.relative_to(Path.cwd())
            saved = save_generated_code(
                result, relative_run_dir, max_lines=160, allowed_outputs=["artifacts/result.json"]
            )
            self.assertEqual(saved, "artifacts/code_pipeline.py")
            self.assertTrue((run_dir / saved).is_file())

    def test_save_generated_code_rejects_invalid_or_framework_owned_output(self):
        with tempfile.TemporaryDirectory() as directory:
            invalid = CodeGenerationResult(
                code="import pd\nPath('agent_trace.md').write_text('bad')",
                entrypoint="artifacts/code_pipeline.py",
            )
            with self.assertRaisesRegex(ValueError, "框架专属文件"):
                save_generated_code(invalid, Path(directory), 160, ["artifacts/result.json"])

            syntax_error = CodeGenerationResult(
                code="def broken(:\n    pass",
                entrypoint="artifacts/code_pipeline.py",
            )
            with self.assertRaisesRegex(ValueError, "编译预检"):
                save_generated_code(syntax_error, Path(directory), 160, ["artifacts/result.json"])

            undeclared_dependency = CodeGenerationResult(
                code="from pypdf import PdfReader",
                entrypoint="artifacts/code_pipeline.py",
            )
            with self.assertRaisesRegex(ValueError, "未声明的 pypdf"):
                save_generated_code(undeclared_dependency, Path(directory), 160, ["artifacts/result.json"])

            unavailable_dependency = CodeGenerationResult(
                code="import framework_test_module_that_is_not_installed",
                entrypoint="artifacts/code_pipeline.py",
            )
            with self.assertRaisesRegex(ValueError, "当前运行环境未安装"):
                save_generated_code(
                    unavailable_dependency,
                    Path(directory),
                    160,
                    ["artifacts/result.json"],
                )


class TraceTests(unittest.TestCase):
    def test_trace_excludes_model_thought_events(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "agent_trace.md"
            messages = [
                SimpleNamespace(source="Agent", type="ThoughtEvent", content="private reasoning"),
                SimpleNamespace(source="Agent", type="TextMessage", content="final result"),
            ]
            save_agent_trace(messages, path)
            trace = path.read_text(encoding="utf-8")
            self.assertNotIn("private reasoning", trace)
            self.assertIn("final result", trace)


class ValidationTests(unittest.TestCase):
    def _spec(self, run_dir: Path):
        return {
            "run_dir": str(run_dir),
            "code_policy": {"mode": "lightweight"},
            "required_artifacts": [
                "artifacts/code_pipeline.py", "artifacts/result.json",
                "task_spec.json", "agent_trace.md", "final_report.md",
            ],
            "required_report_sections": ["任务理解", "结果分析"],
        }

    def test_nonzero_exit_can_only_be_partial_before_report(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            (run_dir / "artifacts").mkdir()
            (run_dir / "artifacts" / "code_pipeline.py").write_text("pass\n", encoding="utf-8")
            (run_dir / "artifacts" / "result.json").write_text("{}", encoding="utf-8")
            (run_dir / "task_spec.json").write_text("{}", encoding="utf-8")
            (run_dir / "agent_trace.md").write_text("trace", encoding="utf-8")
            decision = pre_report_review(self._spec(run_dir), {"execution": {"exit_code": 1}})
            self.assertEqual(decision.status, "partial")
            self.assertTrue(decision.can_generate_final_report)

    def test_report_artifact_reference_stops_before_adjacent_chinese_prose(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            (run_dir / "artifacts").mkdir()
            (run_dir / "artifacts" / "code_pipeline.py").write_text(
                "pass\n", encoding="utf-8",
            )
            (run_dir / "artifacts" / "result.json").write_text(
                "{}", encoding="utf-8",
            )
            (run_dir / "task_spec.json").write_text("{}", encoding="utf-8")
            (run_dir / "agent_trace.md").write_text("trace", encoding="utf-8")
            report = (
                "# 任务理解\n" + "分析内容。" * 80
                + "\n# 结果分析\n所有计数由artifacts/result.json实现独立复算。"
            )
            (run_dir / "final_report.md").write_text(report, encoding="utf-8")

            decision = final_validation(self._spec(run_dir), {
                "execution": {"exit_code": 0},
                "graph_outcome": "success",
                "nodes": {},
                "business_validation": {"status": "passed"},
                "requirement_acceptance": {"status": "passed"},
            })

            self.assertEqual(decision.status, "passed", decision.issues)

    def test_planning_provider_failure_finishes_with_audit_files_not_traceback(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            spec = self._spec(run_dir)
            spec.update({
                "task_id": "provider-failure",
                "task_type": "generic",
                "artifact_contract": {
                    "intermediate_artifacts": ["artifacts/result.json"],
                    "final_artifacts": ["final_report.md"],
                    "framework_artifacts": [
                        "task_spec.json", "agent_trace.md",
                        "run_state.json", "review/validation_report.md",
                    ],
                },
            })
            state = RunState(
                str(run_dir), task_type="generic",
                code_policy=spec["code_policy"],
                required_artifacts=spec["required_artifacts"],
            )
            state.start_stage("classify")
            state.configure(spec)
            state.start_stage("plan")
            decision, _ = _finish_planning_blocked(
                spec, state, [], RuntimeError("provider unavailable"),
            )
            self.assertEqual(decision.status, "failed")
            self.assertEqual(state.get("outcome"), "failed")
            self.assertTrue((run_dir / "agent_trace.md").is_file())
            self.assertTrue((run_dir / "review/validation_report.md").is_file())
            self.assertTrue(any(
                event.get("type") == "planning_recoverable_blocked"
                for event in state.get("events", [])
            ))

    def test_final_report_cannot_be_task_spec(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            (run_dir / "artifacts").mkdir()
            for name, content in {
                "artifacts/code_pipeline.py": "pass\n",
                "artifacts/result.json": "{}",
                "task_spec.json": "{}",
                "agent_trace.md": "trace",
                "final_report.md": "# task_spec.json\n```json\n{}\n```" * 20,
            }.items():
                path = run_dir / name
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(content, encoding="utf-8")
            decision = final_validation(self._spec(run_dir), {
                "execution": {"exit_code": 0},
                "review": {"status": "passed"},
            })
            self.assertEqual(decision.status, "failed")
            self.assertTrue(any("不是业务报告" in issue for issue in decision.issues))

    def test_descriptive_report_policy_rejects_unsupported_decision_claim(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            spec = self._spec(run_dir)
            spec["report_policy"] = {
                "decision_scope": "descriptive_only",
                "forbid_unsupported_decisions": True,
            }
            for name, content in {
                "artifacts/code_pipeline.py": "pass\n",
                "artifacts/result.json": "{}",
                "task_spec.json": "{}",
                "agent_trace.md": "trace",
                "final_report.md": (
                    "# 任务理解\n本报告只应整理事实并开展初步核对。" * 8
                    + "\n## 结果分析\n发现数据异常，建议直接执行删除操作。"
                ),
            }.items():
                path = run_dir / name
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(content, encoding="utf-8")
            decision = final_validation(spec, {
                "execution": {"exit_code": 0},
                "review": {"status": "passed"},
            })
            self.assertEqual(decision.status, "failed")
            self.assertTrue(any("超出初步核对范围" in issue for issue in decision.issues))


class ArtifactValidationTests(unittest.TestCase):
    def _spec(self, run_dir: Path):
        return {
            "run_dir": str(run_dir),
            "artifact_contract": {
                "intermediate_artifacts": ["artifacts/result.json"],
                "final_artifacts": ["final_report.md"],
                "framework_artifacts": [],
            },
        }

    def test_generic_validator_rejects_nonstandard_json(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            (run_dir / "artifacts").mkdir()
            (run_dir / "artifacts" / "result.json").write_text(
                '{"value": NaN}', encoding="utf-8",
            )

            result = validate_intermediate_artifacts(
                self._spec(run_dir), run_dir,
            )

            self.assertEqual(result.status, "failed")
            self.assertEqual(result.failure_type, "artifact_json_invalid")

    def test_generic_validator_accepts_declared_valid_artifact(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            (run_dir / "artifacts").mkdir()
            (run_dir / "artifacts" / "result.json").write_text(
                '{"value": 42}', encoding="utf-8",
            )

            result = validate_intermediate_artifacts(
                self._spec(run_dir), run_dir,
            )

            self.assertEqual(result.status, "passed", result.issues)


class WorkflowIntegrationTests(unittest.IsolatedAsyncioTestCase):
    async def test_business_repair_reuses_persisted_code_as_baseline(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            (run_dir / "artifacts").mkdir()
            marker = "PERSISTED_REPAIR_BASELINE"
            (run_dir / "artifacts/code_pipeline.py").write_text(
                f"# {marker}\n", encoding="utf-8",
            )
            spec = {
                "task_id": "business-repair-baseline",
                "task_name": "generic repair",
                "task_type": "general",
                "run_dir": str(run_dir),
                "code_policy": {
                    "mode": "lightweight", "max_lines": 80,
                    "max_retries": 0,
                },
                "artifact_contract": {
                    "intermediate_artifacts": [
                        "artifacts/code_pipeline.py",
                        "artifacts/result.json",
                    ],
                    "final_artifacts": [], "framework_artifacts": [],
                },
                "artifacts": {"code": "artifacts/code_pipeline.py"},
            }
            generated = {
                "language": "python",
                "code": (
                    "import json\nfrom pathlib import Path\n"
                    "Path('artifacts/result.json').write_text("
                    "json.dumps({'value': 2}), encoding='utf-8')"
                ),
                "entrypoint": "artifacts/code_pipeline.py",
                "expected_outputs": ["artifacts/result.json"],
            }
            state = RunState(
                str(run_dir), task_type="general",
                code_policy=spec["code_policy"],
            )
            state.configure(spec)
            state.start_stage("execute")
            executor = CodePipelineNodeExecutor(
                ReplayChatCompletionClient([]), state, [], {},
            )
            node = TaskNode(
                node_id="code", description="repair existing implementation",
                capability="code",
                output_artifacts=[
                    "artifacts/code_pipeline.py", "artifacts/result.json",
                ],
                recovery_context={
                    "kind": "post_business_validation_repair",
                    "repair_instruction": "repair verified field mismatch",
                },
                success_criteria=[{"type": "execution_exit_code", "equals": 0}],
            )

            mocked = AsyncMock(return_value=json.dumps(generated))
            with patch(
                "orchestration.execution.default_executors.run_agent", mocked,
            ):
                result = await executor.execute(NodeExecutionContext(
                    graph_id="g", node=node, task_spec=spec,
                    run_dir=str(run_dir),
                ))

            self.assertEqual(result.status, "completed", result.error_message)
            prompt = mocked.await_args.args[1]
            self.assertIn(marker, prompt)
            self.assertIn("__previous_code_candidate__", prompt)

    async def test_static_code_repair_does_not_consume_business_retry(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            (run_dir / "artifacts").mkdir()
            spec = {
                "task_id": "format-retry",
                "task_name": "generic code generation",
                "task_type": "general",
                "run_dir": str(run_dir),
                "code_policy": {
                    "mode": "lightweight", "max_lines": 80,
                    "max_retries": 0,
                },
                "artifact_contract": {
                    "intermediate_artifacts": [
                        "artifacts/code_pipeline.py",
                        "artifacts/result.json",
                    ],
                    "final_artifacts": [], "framework_artifacts": [],
                },
                "artifacts": {"code": "artifacts/code_pipeline.py"},
            }
            invalid = {
                "language": "python", "code": "value = '",
                "entrypoint": "artifacts/code_pipeline.py",
                "expected_outputs": ["artifacts/result.json"],
            }
            valid = {
                "language": "python",
                "code": (
                    "import json\nfrom pathlib import Path\n"
                    "Path('artifacts/result.json').write_text("
                    "json.dumps({'value': 1}), encoding='utf-8')"
                ),
                "entrypoint": "artifacts/code_pipeline.py",
                "expected_outputs": ["artifacts/result.json"],
            }
            state = RunState(
                str(run_dir), task_type="general",
                code_policy=spec["code_policy"],
            )
            state.configure(spec)
            state.start_stage("execute")
            executor = CodePipelineNodeExecutor(
                ReplayChatCompletionClient([
                    json.dumps(invalid), json.dumps(valid),
                ]),
                state, [], {},
            )
            node = TaskNode(
                node_id="code", description="implement fixed specification",
                capability="code",
                output_artifacts=[
                    "artifacts/code_pipeline.py", "artifacts/result.json",
                ],
                success_criteria=[{"type": "execution_exit_code", "equals": 0}],
            )

            result = await executor.execute(NodeExecutionContext(
                graph_id="g", node=node, task_spec=spec,
                run_dir=str(run_dir),
            ))

            self.assertEqual(result.status, "completed", result.error_message)
            self.assertEqual(state.get("execution")["attempts"], 2)
            events = [event["type"] for event in state.get("events", [])]
            self.assertIn("code_generation_format_retry_scheduled", events)
            format_event = next(
                event for event in state.get("events", [])
                if event["type"] == "code_generation_format_retry_scheduled"
            )
            self.assertFalse(
                format_event["payload"]["business_retry_consumed"],
            )

    async def test_explicit_workflow_reaches_success_and_reports_once(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            (run_dir / "artifacts").mkdir()
            spec = {
                "task_id": "test_run",
                "task_name": "测试任务",
                "task_type": "document_analysis",
                "run_dir": str(run_dir),
                "input": {"type": "directory", "files": [], "text": ""},
                "file_previews": [],
                "code_policy": {
                    "mode": "lightweight", "allows_complex": False,
                    "allows_lightweight": True, "max_lines": 80, "max_retries": 1,
                },
                "required_artifacts": [
                    "artifacts/result.json", "artifacts/code_pipeline.py",
                    "task_spec.json", "agent_trace.md", "final_report.md",
                ],
                "required_report_sections": ["任务理解", "输入文件清单", "分析过程", "结果分析", "风险与改进建议", "最终交付物"],
                "success_criteria": {"required_artifacts": [], "required_report_sections": []},
                "artifacts": {
                    "code": "artifacts/code_pipeline.py",
                    "results": ["artifacts/result.json"],
                    "report": "final_report.md", "trace": "agent_trace.md",
                },
            }
            (run_dir / "task_spec.json").write_text(json.dumps(spec), encoding="utf-8")
            state = RunState(
                str(run_dir), task_type=spec["task_type"], code_policy=spec["code_policy"],
                required_artifacts=spec["required_artifacts"],
            )
            state.start_stage("classify")
            state.record_artifact("task_spec.json")

            plan = {
                "graph_id": "test_run_graph", "version": 1, "goal": "生成并验证结构化结果",
                "nodes": [
                    {
                        "node_id": "generate_result", "description": "生成并执行结果代码",
                        "capability": "code", "dependencies": [], "input_artifacts": [],
                        "output_artifacts": ["artifacts/result.json", "artifacts/code_pipeline.py"],
                        "success_criteria": [{"type": "execution_exit_code", "equals": 0}],
                        "max_retries": 0,
                    },
                    {
                        "node_id": "validate_result", "description": "确定性校验中间产物",
                        "capability": "artifact_validation", "dependencies": ["generate_result"],
                        "input_artifacts": ["artifacts/result.json"], "output_artifacts": [],
                        "success_criteria": [{"type": "artifact_quality", "equals": "passed"}],
                        "max_retries": 0,
                    },
                ],
            }
            code = {
                "language": "python",
                "code": "from pathlib import Path\nPath('artifacts/result.json').write_text('{\\\"value\\\": 1}', encoding='utf-8')",
                "entrypoint": "artifacts/code_pipeline.py",
                "expected_outputs": ["artifacts/result.json"], "notes": "",
            }
            review = {
                "blocking_issues": [],
                "advisory_issues": ["建议人工复核并补充说明"],
                "evidence_references": ["artifacts/result.json:value"],
                "repair_recommended": False,
                "repair_target": "none",
            }
            report = """# 任务理解
这是一次用于验证显式工作流的业务测试，内容完全来源于框架提供的证据。
## 输入文件清单
本测试没有外部输入文件，使用内置的确定性测试数据完成验证。
## 分析过程
流程生成可复现脚本，执行后产生结构化 JSON，再进行报告前审查。
## 结果分析
结构化结果中的 value 为 1，脚本正常退出，所有必需产物均可读取。
## 风险与改进建议
生产环境仍应为模型超时、网络异常和无效 JSON 配置重试与监控。
## 最终交付物
本轮交付结构化结果、可复现脚本和业务报告，所有文件均已通过存在性检查。
"""
            client = ReplayChatCompletionClient([
                json.dumps(plan, ensure_ascii=False),
                json.dumps(code, ensure_ascii=False),
                json.dumps(review, ensure_ascii=False),
                report,
            ])
            decision, _ = await run_explicit_workflow(spec, state, client)
            self.assertEqual(decision.status, "passed", decision.issues)
            self.assertEqual(state.get("outcome"), "success")
            self.assertTrue((run_dir / "artifacts" / "code_pipeline.py").is_file())
            self.assertTrue((run_dir / "artifacts" / "result.json").is_file())
            self.assertEqual(
                state.get("execution")["history"][0]["produced_files"],
                ["artifacts/code_pipeline.py", "artifacts/result.json"],
            )
            self.assertTrue((run_dir / "final_report.md").read_text(encoding="utf-8").startswith("# 任务理解"))

    async def test_artifact_quality_failure_routes_back_to_code(self):
        with tempfile.TemporaryDirectory() as directory:
            run_dir = Path(directory)
            (run_dir / "artifacts").mkdir()
            spec = {
                "task_id": "retry_run", "task_name": "通用产物重试测试",
                "task_type": "document_analysis",
                "run_dir": str(run_dir),
                "input": {"type": "directory", "files": [], "text": ""}, "file_previews": [],
                "code_policy": {"mode": "lightweight", "max_lines": 160, "max_retries": 1},
                "artifact_contract": {
                    "intermediate_artifacts": ["artifacts/result.json", "artifacts/code_pipeline.py"],
                    "final_artifacts": ["final_report.md"],
                    "framework_artifacts": ["task_spec.json", "agent_trace.md", "run_state.json"],
                },
                "required_artifacts": [
                    "artifacts/result.json", "artifacts/code_pipeline.py",
                    "final_report.md", "task_spec.json", "agent_trace.md", "run_state.json",
                ],
                "required_report_sections": ["任务理解", "结果分析"],
                "artifacts": {"code": "artifacts/code_pipeline.py", "results": ["artifacts/result.json"]},
            }
            (run_dir / "task_spec.json").write_text(json.dumps(spec), encoding="utf-8")
            state = RunState(
                str(run_dir), task_type=spec["task_type"], code_policy=spec["code_policy"],
                required_artifacts=spec["required_artifacts"],
            )
            state.start_stage("classify")
            state.configure(spec)

            plan = {
                "graph_id": "retry_run_graph", "version": 1, "goal": "生成并验证通用 JSON 产物",
                "nodes": [
                    {
                        "node_id": "generate_result", "description": "生成并执行数据处理代码",
                        "capability": "code", "dependencies": [], "input_artifacts": [],
                        "output_artifacts": [
                            "artifacts/result.json", "artifacts/code_pipeline.py"
                        ],
                        "success_criteria": [{"type": "execution_exit_code", "equals": 0}],
                        "max_retries": 0,
                    },
                    {
                        "node_id": "validate_result", "description": "确定性校验 JSON 产物",
                        "capability": "artifact_validation", "dependencies": ["generate_result"],
                        "input_artifacts": ["artifacts/result.json"], "output_artifacts": [],
                        "success_criteria": [{"type": "artifact_quality", "equals": "passed"}],
                        "max_retries": 0,
                    },
                ],
            }
            invalid_code = {
                "language": "python",
                "code": "from pathlib import Path\nPath('artifacts/result.json').write_text('{\\\"value\\\": NaN}')",
                "entrypoint": "artifacts/code_pipeline.py",
                "expected_outputs": ["artifacts/result.json"], "notes": "first attempt",
            }
            valid_artifact = {"value": 42, "status": "validated"}
            valid_code_text = (
                "import json\nfrom pathlib import Path\n"
                f"data = {valid_artifact!r}\n"
                "Path('artifacts/result.json').write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')"
            )
            valid_code = {
                "language": "python", "code": valid_code_text,
                "entrypoint": "artifacts/code_pipeline.py",
                "expected_outputs": ["artifacts/result.json"], "notes": "repaired",
            }
            semantic_review = {
                "blocking_issues": [], "advisory_issues": [], "evidence_references": [],
                "repair_recommended": False, "repair_target": "none",
            }
            report = (
                "# 任务理解\n本报告验证产物失败后能返回代码阶段完成一次局部修复。\n"
                "框架使用固定状态机推进，所有结论均来自通过契约校验的中间产物。\n"
                "## 结果分析\n修复后的结构化结果 value 为42，状态为 validated。\n"
                "代码真实执行成功，产物通过存在性与标准 JSON 检查。\n"
                "该验证说明失败被精确路由到 execute 阶段，没有重新执行无关的准备与规划工作。\n"
                "最终报告仅消费经过框架验证的数据，因此不会把第一次失败产生的非法数值带入交付结果。\n"
            )
            client = ReplayChatCompletionClient([
                json.dumps(plan, ensure_ascii=False),
                json.dumps(invalid_code, ensure_ascii=False),
                json.dumps(valid_code, ensure_ascii=False),
                json.dumps(semantic_review, ensure_ascii=False),
                report,
            ])
            decision, _ = await run_explicit_workflow(spec, state, client)
            self.assertEqual(decision.status, "passed", decision.issues)
            current = state.to_dict()
            self.assertEqual(current["execution"]["attempts"], 2)
            self.assertTrue(all(
                set(item["produced_files"]) <= {
                    "artifacts/code_pipeline.py", "artifacts/result.json",
                }
                for item in current["execution"]["history"]
            ))
            self.assertEqual([item["status"] for item in current["artifact_quality"]["history"]], ["failed", "passed"])
            self.assertIsNone(current["failure"])
            event_types = [event["type"] for event in current["events"]]
            self.assertIn("retry_scheduled", event_types)
            self.assertIn("failure_cleared", event_types)


if __name__ == "__main__":
    unittest.main()
