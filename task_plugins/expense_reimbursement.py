from __future__ import annotations

from datetime import date, datetime
import math
from pathlib import Path
import re
from typing import Any

import pandas as pd
import pdfplumber
from docx import Document
from pydantic import BaseModel, Field, ValidationError, field_validator

from orchestration.schemas import ArtifactQualityResult
from task_plugins.base import load_strict_json
from task_plugins.generic import GenericTaskPlugin


class ExpenseRecord(BaseModel):
    date: date
    category: str
    amount: float
    invoice_status: str
    note: str = ""
    issues: list[str] = Field(default_factory=list)

    @field_validator("category", "invoice_status")
    @classmethod
    def non_empty_text(cls, value: str):
        if not value.strip() or value.strip().lower() in {"nan", "nat", "none"}:
            raise ValueError("不能为空或 nan/NaT")
        return value.strip()

    @field_validator("amount")
    @classmethod
    def finite_amount(cls, value: float):
        if not math.isfinite(value):
            raise ValueError("金额必须是有限数字")
        return value

    @field_validator("issues")
    @classmethod
    def issues_only(cls, values: list[str]):
        if any(value.strip().lower() in {"ok", "正常", "通过"} for value in values):
            raise ValueError("issues 只能记录问题，不能写 OK/正常")
        return values


class ExpenseTotals(BaseModel):
    total_amount: float
    by_category: dict[str, float]
    missing_invoice_count: int = 0

    @field_validator("total_amount")
    @classmethod
    def finite_total(cls, value: float):
        if not math.isfinite(value):
            raise ValueError("total_amount 必须是有限数字")
        return value

    @field_validator("by_category")
    @classmethod
    def finite_categories(cls, values: dict[str, float]):
        if not values or any(not key.strip() or not math.isfinite(value) for key, value in values.items()):
            raise ValueError("分类汇总必须非空且金额有限")
        return values


class ExpenseSummaryArtifact(BaseModel):
    trip_info: dict = Field(default_factory=dict)
    policy_rules: dict
    expense_records: list[ExpenseRecord]
    summary: ExpenseTotals
    issues: list[str] = Field(default_factory=list)

    @field_validator("policy_rules")
    @classmethod
    def policy_must_be_present(cls, value: dict):
        if not value:
            raise ValueError("policy_rules 不能为空")
        return value

    @field_validator("expense_records")
    @classmethod
    def records_must_be_present(cls, value: list[ExpenseRecord]):
        if not value:
            raise ValueError("expense_records 不能为空")
        return value


class ExpenseReimbursementPlugin(GenericTaskPlugin):
    plugin_id = "expense_reimbursement"

    @staticmethod
    def _json_value(value: Any):
        if value is None:
            return None
        try:
            if pd.isna(value):
                return None
        except (TypeError, ValueError):
            pass
        if isinstance(value, pd.Timestamp):
            return value.date().isoformat()
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if hasattr(value, "item"):
            value = value.item()
        if isinstance(value, float) and not math.isfinite(value):
            return None
        return value

    def prepare_inputs(self, task_spec: dict, run_dir: Path) -> dict:
        """Normalize heterogeneous expense inputs before any LLM-written code.

        The source spreadsheet puts detail columns and a reference summary next
        to each other. They are intentionally extracted as two independent
        datasets so a generated script cannot mistake a populated summary cell
        for a summary *row* and discard valid details.
        """
        prepared: dict[str, Any] = {
            "expense_rows": [],
            "provided_summary": [],
            "travel_application": {"paragraphs": [], "tables": []},
            "travel_facts": {},
            "policy_text": "",
            "policy_rules": {},
        }
        input_dir = run_dir / "inputs"
        if not input_dir.is_dir():
            return prepared
        for path in sorted(input_dir.iterdir()):
            suffix = path.suffix.lower()
            if suffix in {".xlsx", ".xls"}:
                sheets = pd.read_excel(path, sheet_name=None)
                for sheet_name, frame in sheets.items():
                    detail_columns = [name for name in ["日期", "费用类型", "金额", "发票状态", "备注"] if name in frame]
                    for _, row in frame.iterrows():
                        detail = {name: self._json_value(row.get(name)) for name in detail_columns}
                        identity = [detail.get(name) for name in ["日期", "费用类型", "金额"] if name in detail]
                        if identity and any(value not in (None, "") for value in identity):
                            detail["sheet"] = sheet_name
                            prepared["expense_rows"].append(detail)
                    summary_name = next((name for name in ["汇总项", "汇总", "总计"] if name in frame), None)
                    summary_value = next((name for name in ["金额/数量", "汇总金额", "数量"] if name in frame), None)
                    if summary_name:
                        for _, row in frame.iterrows():
                            name = self._json_value(row.get(summary_name))
                            if name not in (None, ""):
                                prepared["provided_summary"].append({
                                    "item": name,
                                    "value": self._json_value(row.get(summary_value)) if summary_value else None,
                                    "sheet": sheet_name,
                                })
            elif suffix == ".docx":
                document = Document(path)
                prepared["travel_application"]["paragraphs"] = [
                    paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()
                ]
                prepared["travel_application"]["tables"] = [
                    [[cell.text.strip() for cell in row.cells] for row in table.rows]
                    for table in document.tables
                ]
            elif suffix == ".pdf":
                with pdfplumber.open(path) as pdf:
                    prepared["policy_text"] = "\n".join(page.extract_text() or "" for page in pdf.pages).strip()
        prepared["travel_facts"] = self._extract_travel_facts(
            prepared["travel_application"]
        )
        prepared["policy_rules"] = self._extract_policy_rules(prepared["policy_text"])
        return prepared

    @staticmethod
    def _extract_travel_facts(travel_application: dict) -> dict[str, str]:
        """Turn formatting-sensitive Word content into a stable fact boundary."""
        paragraphs = [str(value) for value in travel_application.get("paragraphs", [])]
        table_pairs: dict[str, str] = {}
        for table in travel_application.get("tables", []):
            for row in table:
                if isinstance(row, list) and len(row) >= 2:
                    table_pairs[str(row[0]).strip()] = str(row[1]).strip()

        application_id = ""
        applicant = table_pairs.get("员工姓名", "")
        for paragraph in paragraphs:
            if match := re.search(r"申请编号[：:]\s*(.+)", paragraph):
                application_id = match.group(1).strip()
            if match := re.search(r"申请人[：:]\s*(.+)", paragraph):
                applicant = match.group(1).strip()
        dates = re.findall(r"\d{4}-\d{2}-\d{2}", table_pairs.get("出差时间", ""))
        return {
            "申请编号": application_id or "缺失",
            "申请人": applicant or "缺失",
            "出差开始": dates[0] if dates else "缺失",
            "出差结束": dates[1] if len(dates) > 1 else "缺失",
        }

    @staticmethod
    def _extract_policy_rules(policy_text: str) -> dict[str, str]:
        """Extract stable policy snippets so generated code need not parse PDF layout."""
        compact = re.sub(r"\s+", " ", policy_text).strip()
        patterns = {
            "住宿标准": r"一线城市住宿费不超过\s*\d+\s*元/晚；其他城市住宿费不超过\s*\d+\s*元/晚",
            "交通标准": r"高铁二等座、动车二等座、飞机经济舱可报销",
            "餐补标准": r"餐补标准为每人每天\s*\d+\s*元",
            "发票要求": r"所有报销项目必须提供合规发票",
        }
        return {
            name: (match.group(0).strip() if (match := re.search(pattern, compact)) else "缺失")
            for name, pattern in patterns.items()
        }

    def code_instructions(self, task_spec: dict) -> str:
        return """
必须只读取框架生成的 normalized_input.json，不要重新读取 inputs/ 中的 Word/PDF/Excel。
normalized_input.json 中 expense_rows 是已确定的真实明细；provided_summary 只是原文件的
参考汇总，两者横向并存，不得因为 provided_summary 非空而过滤 expense_rows。
框架已经确定性处理 Word/PDF 的全角标点、表格形状与断行：输出 JSON 的 trip_info 必须
直接复制 normalized_input.travel_facts，policy_rules 必须直接复制
normalized_input.policy_rules。禁止再次从 paragraphs、tables 或 policy_text 猜测这些字段。
超出出差日期及未提供发票的明细必须在该条 issues 中明确标记。
expense_summary.json 必须严格采用以下结构：
{
  "trip_info": {"申请编号": "...", "申请人": "...", "出差开始": "YYYY-MM-DD或缺失", "出差结束": "YYYY-MM-DD或缺失"},
  "policy_rules": {"住宿标准": "...", "交通标准": "...", "餐补标准": "...", "发票要求": "..."},
  "expense_records": [
    {"date": "YYYY-MM-DD", "category": "住宿", "amount": 480.0,
     "invoice_status": "已提供", "note": "...", "issues": []}
  ],
  "summary": {"total_amount": 1485.0, "by_category": {"住宿": 960.0}, "missing_invoice_count": 1},
  "issues": ["可追溯的问题说明"]
}
必须过滤日期、类别和金额同时为空的行；禁止 NaN/NaT；issues 不得写 OK。
summary.total_amount 必须等于明细金额之和，by_category 之和也必须等于 total_amount。
只输出 artifacts/expense_summary.json，报告和框架文件由框架生成。
""".strip()

    def validate_intermediate(self, task_spec: dict, run_dir: Path) -> ArtifactQualityResult:
        generic = super().validate_intermediate(task_spec, run_dir)
        if generic.status == "failed":
            return generic

        relative = "artifacts/expense_summary.json"
        path = run_dir / relative
        if not path.is_file():
            return ArtifactQualityResult(
                status="failed", checked_artifacts=[relative], issues=[f"缺少中间产物：{relative}"],
                failure_type="artifact_missing", repair_target="code", resume_stage="execute",
            )
        try:
            artifact = ExpenseSummaryArtifact.model_validate(load_strict_json(path))
        except (OSError, ValueError, ValidationError) as exc:
            return ArtifactQualityResult(
                status="failed", checked_artifacts=[relative],
                issues=[f"报销产物 Schema 校验失败：{exc}"],
                failure_type="artifact_schema_failure", repair_target="code", resume_stage="execute",
            )

        issues = []
        record_total = sum(record.amount for record in artifact.expense_records)
        category_total = sum(artifact.summary.by_category.values())
        if not math.isclose(record_total, artifact.summary.total_amount, rel_tol=1e-9, abs_tol=0.01):
            issues.append(f"明细合计 {record_total} 与 total_amount {artifact.summary.total_amount} 不一致")
        if not math.isclose(category_total, artifact.summary.total_amount, rel_tol=1e-9, abs_tol=0.01):
            issues.append(f"分类合计 {category_total} 与 total_amount {artifact.summary.total_amount} 不一致")
        missing = sum("未提供" in record.invoice_status for record in artifact.expense_records)
        if missing != artifact.summary.missing_invoice_count:
            issues.append(
                f"缺票明细 {missing} 条，与 missing_invoice_count {artifact.summary.missing_invoice_count} 不一致"
            )

        normalized_path = run_dir / "normalized_input.json"
        if normalized_path.is_file():
            try:
                normalized = load_strict_json(normalized_path)
                issues.extend(self._validate_against_normalized(artifact, normalized))
            except (OSError, ValueError) as exc:
                issues.append(f"框架标准化输入无法解析：{exc}")
        return ArtifactQualityResult(
            status="failed" if issues else "passed",
            checked_artifacts=[relative],
            issues=issues,
            failure_type="artifact_schema_failure" if issues else "none",
            repair_target="code" if issues else "none",
            resume_stage="execute" if issues else "none",
        )

    @staticmethod
    def _validate_against_normalized(artifact: ExpenseSummaryArtifact, normalized: dict) -> list[str]:
        """Cross-check business output against deterministic normalized evidence."""
        issues: list[str] = []
        travel = normalized.get("travel_application", {})
        paragraphs = [str(value) for value in travel.get("paragraphs", [])]
        table_pairs: dict[str, str] = {}
        for table in travel.get("tables", []):
            for row in table:
                if isinstance(row, list) and len(row) >= 2:
                    table_pairs[str(row[0]).strip()] = str(row[1]).strip()

        expected_id = ""
        expected_applicant = table_pairs.get("员工姓名", "")
        for paragraph in paragraphs:
            if match := re.search(r"申请编号[：:]\s*(.+)", paragraph):
                expected_id = match.group(1).strip()
            if match := re.search(r"申请人[：:]\s*(.+)", paragraph):
                expected_applicant = match.group(1).strip()
        trip_dates = re.findall(r"\d{4}-\d{2}-\d{2}", table_pairs.get("出差时间", ""))
        expected_trip = {
            "申请编号": expected_id,
            "申请人": expected_applicant,
            "出差开始": trip_dates[0] if trip_dates else "",
            "出差结束": trip_dates[1] if len(trip_dates) > 1 else "",
        }
        for field, expected in expected_trip.items():
            actual = str(artifact.trip_info.get(field, "")).strip()
            if expected and actual != expected:
                issues.append(f"trip_info.{field} 应为 {expected!r}，实际为 {actual!r}")

        source_rows = normalized.get("expense_rows", [])
        if source_rows and len(artifact.expense_records) != len(source_rows):
            issues.append(
                f"报销明细数量应为 {len(source_rows)}，实际为 {len(artifact.expense_records)}"
            )
        output_by_key = {
            (record.date.isoformat(), record.category, round(record.amount, 2)): record
            for record in artifact.expense_records
        }
        start = date.fromisoformat(expected_trip["出差开始"]) if expected_trip["出差开始"] else None
        end = date.fromisoformat(expected_trip["出差结束"]) if expected_trip["出差结束"] else None
        for row in source_rows:
            try:
                key = (str(row.get("日期", ""))[:10], str(row.get("费用类型", "")).strip(), round(float(row["金额"]), 2))
            except (KeyError, TypeError, ValueError):
                continue
            record = output_by_key.get(key)
            if record is None:
                issues.append(f"缺少源明细映射：日期={key[0]}，类别={key[1]}，金额={key[2]}")
                continue
            issue_text = " ".join(record.issues)
            if "未提供" in str(row.get("发票状态", "")) and not any(word in issue_text for word in ["发票", "缺票"]):
                issues.append(f"{key[0]} {key[1]} 未提供发票，但明细 issues 未标记")
            row_date = date.fromisoformat(key[0])
            if start and end and not start <= row_date <= end and not any(word in issue_text for word in ["超出", "日期"]):
                issues.append(f"{key[0]} {key[1]} 超出出差日期，但明细 issues 未标记")

        if normalized.get("policy_text"):
            vague = [
                name for name, value in artifact.policy_rules.items()
                if not str(value).strip() or "见文档" in str(value)
            ]
            if vague:
                issues.append(f"政策原文完整，但以下规则未具体提取：{', '.join(vague)}")
        return issues
